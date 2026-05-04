import streamlit as st
import streamlit.components.v1 as components
import tempfile
import os
import difflib
import html
from datetime import datetime
from docx import Document

from modules.word_compare.comparator import compare_documents


# --------------------------------------------------
# Extract document content
# --------------------------------------------------
def extract_doc_content(doc_file):
    doc = Document(doc_file)
    content = []

    # Paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content.append(text)

    # Tables
    for table_index, table in enumerate(doc.tables):
        content.append(f"[TABLE-{table_index+1}]")

        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip()
                for cell in row.cells
            )
            content.append(f"[TABLE] {row_text}")

    # Images
    image_count = 0
    for rel in doc.part.rels.values():
        try:
            if rel.is_external:
                continue

            if "image" in rel.target_ref.lower():
                image_count += 1
        except Exception:
            continue

    if image_count:
        content.append(
            f"[IMAGES FOUND: {image_count}]"
        )

    return content


# --------------------------------------------------
# Create aligned rows for both previews
# --------------------------------------------------
def generate_aligned_diff_rows(old_lines, new_lines):
    matcher = difflib.SequenceMatcher(
        None,
        old_lines,
        new_lines
    )

    old_rows = []
    new_rows = []

    for opcode, i1, i2, j1, j2 in matcher.get_opcodes():

        # ---------------------------
        # Equal rows
        # ---------------------------
        if opcode == "equal":
            for old_line, new_line in zip(
                old_lines[i1:i2],
                new_lines[j1:j2]
            ):
                old_rows.append(
                    build_row(old_line, "normal")
                )
                new_rows.append(
                    build_row(new_line, "normal")
                )

        # ---------------------------
        # Delete rows
        # ---------------------------
        elif opcode == "delete":
            deleted_lines = old_lines[i1:i2]

            for line in deleted_lines:
                old_rows.append(
                    build_row(
                        f"❌ Removed: {line}",
                        "removed"
                    )
                )

                # placeholder in new
                new_rows.append(
                    build_row("", "blank")
                )

        # ---------------------------
        # Insert rows
        # ---------------------------
        elif opcode == "insert":
            inserted_lines = new_lines[j1:j2]

            for line in inserted_lines:
                old_rows.append(
                    build_row("", "blank")
                )

                new_rows.append(
                    build_row(
                        f"➕ Added: {line}",
                        "added"
                    )
                )

        # ---------------------------
        # Replace rows
        # ---------------------------
        elif opcode == "replace":
            old_chunk = old_lines[i1:i2]
            new_chunk = new_lines[j1:j2]

            max_len = max(
                len(old_chunk),
                len(new_chunk)
            )

            for idx in range(max_len):
                old_line = (
                    old_chunk[idx]
                    if idx < len(old_chunk)
                    else ""
                )

                new_line = (
                    new_chunk[idx]
                    if idx < len(new_chunk)
                    else ""
                )

                old_rows.append(
                    build_row(
                        f"🔄 Replaced: {old_line}"
                        if old_line else "",
                        "updated" if old_line else "blank"
                    )
                )

                new_rows.append(
                    build_row(
                        f"🔄 Updated: {new_line}"
                        if new_line else "",
                        "updated" if new_line else "blank"
                    )
                )

    return "".join(old_rows), "".join(new_rows)


# --------------------------------------------------
# Build row html
# --------------------------------------------------
def build_row(text, css_class):
    safe_text = html.escape(text)

    return f"""
    <div 
        class="line {css_class}"
        title="{safe_text}"
    >
        {safe_text}
    </div>
    """


# --------------------------------------------------
# Render synchronized preview
# --------------------------------------------------
def render_synced_preview(
    old_html,
    new_html
):
    combined_html = f"""
    <html>
    <head>
    <style>
        body {{
            margin:0;
            font-family:Arial;
        }}

        .container {{
            display:flex;
            width:100%;
            height:650px;
            border:1px solid #ccc;
        }}

        .pane {{
            width:50%;
            overflow-y:auto;
            border-right:1px solid #ddd;
        }}

        .line {{
            height:32px;
            line-height:20px;
            padding:6px;
            margin:2px;
            border-radius:4px;
            font-size:13px;
            white-space:nowrap;
            overflow:hidden;
            text-overflow:ellipsis;
            box-sizing:border-box;
        }}

        .normal {{
            background:white;
        }}

        .removed {{
            background:#ffcccc;
        }}

        .added {{
            background:#ccffcc;
        }}

        .updated {{
            background:#ffe599;
        }}

        .blank {{
            background:white;
        }}
    </style>
    </head>

    <body>

    <div class="container">

        <div class="pane" id="leftPane">
            {old_html}
        </div>

        <div class="pane" id="rightPane">
            {new_html}
        </div>

    </div>

    <script>
        const left = document.getElementById("leftPane");
        const right = document.getElementById("rightPane");

        let syncing = false;

        left.addEventListener("scroll", function() {{
            if (!syncing) {{
                syncing = true;
                right.scrollTop = left.scrollTop;
                syncing = false;
            }}
        }});

        right.addEventListener("scroll", function() {{
            if (!syncing) {{
                syncing = true;
                left.scrollTop = right.scrollTop;
                syncing = false;
            }}
        }});
    </script>

    </body>
    </html>
    """

    components.html(
        combined_html,
        height=700,
        scrolling=False
    )


# --------------------------------------------------
# Main UI
# --------------------------------------------------
def render():
    st.title("Word Compare Utility")

    st.write(
        "Upload old master document and new document "
        "to preview and highlight changes."
    )

    col1, col2 = st.columns(2)

    with col1:
        old_file = st.file_uploader(
            "Upload Old Document (Master)",
            type=["docx"]
        )

    with col2:
        new_file = st.file_uploader(
            "Upload New Document",
            type=["docx"]
        )

    if old_file and new_file:

        old_file.seek(0)
        old_lines = extract_doc_content(old_file)

        new_file.seek(0)
        new_lines = extract_doc_content(new_file)

        old_html, new_html = generate_aligned_diff_rows(
            old_lines,
            new_lines
        )

        st.subheader("Difference Preview")

        # filenames OUTSIDE preview pane
        h1, h2 = st.columns(2)

        with h1:
            st.markdown(
                f"### {old_file.name}"
            )

        with h2:
            st.markdown(
                f"### {new_file.name}"
            )

        render_synced_preview(
            old_html,
            new_html
        )

        st.divider()

        if st.button(
            "Generate Highlighted Word File"
        ):
            try:
                old_file.seek(0)
                new_file.seek(0)

                base_name = os.path.splitext(
                    old_file.name
                )[0]

                current_date = datetime.now().strftime(
                    "%d%b%Y"
                )

                output_filename = (
                    f"{base_name}"
                    f"_Diff-Highlighted_"
                    f"{current_date}.docx"
                )

                with tempfile.NamedTemporaryFile(
                    delete=False,
                    suffix=".docx"
                ) as tmp:
                    output_path = tmp.name

                compare_documents(
                    old_file,
                    new_file,
                    output_path
                )

                with open(output_path, "rb") as f:
                    file_bytes = f.read()

                st.download_button(
                    label="Download Compared Document",
                    data=file_bytes,
                    file_name=output_filename,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

                os.remove(output_path)

            except Exception as e:
                st.error(
                    f"Error: {str(e)}"
                )
