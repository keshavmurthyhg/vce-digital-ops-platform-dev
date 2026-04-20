from docx import Document
from docx.shared import Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from io import BytesIO
import re


def clean_text(text):
    return re.sub(r"How does the user want.*?\d+", "", str(text), flags=re.I).strip()


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    text_elem = OxmlElement("w:t")
    text_elem.text = text

    run.append(text_elem)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def generate_word_doc(data, root, l2, res, images=None):

    doc = Document()

    title = doc.add_heading("INCIDENT REPORT", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True
    title.runs[0].font.color.rgb = RGBColor(31, 78, 121)

    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"

    def fill(r, c, key, val):
        h = table.rows[r].cells[c]
        v = table.rows[r].cells[c+1]

        h.text = key.upper()

        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "D9D9D9")
        h._element.get_or_add_tcPr().append(shade)

        for run in h.paragraphs[0].runs:
            run.bold = True

        p = v.paragraphs[0]

        if key == "Incident":
            add_hyperlink(p, f"https://volvoitsm.service-now.com/nav_to.do?uri=incident.do?sysparm_query=number={val}", str(val))
        elif key == "Azure Bug":
            add_hyperlink(p, f"https://dev.azure.com/VolvoGroup-DVP/VCEWindchillPLM/_workitems/edit/{val}", str(val))
        elif key == "PTC Case":
            add_hyperlink(p, f"https://support.ptc.com/app/caseviewer/?case={val}", str(val))
        else:
            p.text = str(val)

    fill(0,0,"Incident",data.get("number"))
    fill(0,2,"Created By",data.get("created_by"))
    fill(1,0,"Azure Bug",data.get("azure_bug"))
    fill(1,2,"Created Date",data.get("created_date"))
    fill(2,0,"PTC Case",data.get("ptc_case"))
    fill(2,2,"Assigned To",data.get("assigned_to"))
    fill(3,0,"Priority",data.get("priority"))
    fill(3,2,"Resolved Date",data.get("resolved_date"))

    doc.add_paragraph("")

    t2 = doc.add_table(rows=2, cols=2)
    t2.style = "Table Grid"

    t2.rows[0].cells[0].text = "SHORT DESCRIPTION"
    t2.rows[0].cells[1].text = "DESCRIPTION"

    t2.rows[1].cells[0].text = clean_text(data.get("short_description"))
    t2.rows[1].cells[1].text = clean_text(data.get("description"))

    doc.add_heading("ROOT CAUSE", 1)
    doc.add_paragraph(root)

    if images and images.get("root"):
        doc.add_picture(images["root"], width=Inches(5))

    doc.add_heading("L2 ANALYSIS", 1)
    doc.add_paragraph(l2)

    if images and images.get("l2"):
        doc.add_picture(images["l2"], width=Inches(5))

    doc.add_heading("RESOLUTION", 1)
    doc.add_paragraph(res)

    if images and images.get("res"):
        doc.add_picture(images["res"], width=Inches(5))

    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.text = f"{data.get('number')}    Page    {data.get('priority')}"

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer
